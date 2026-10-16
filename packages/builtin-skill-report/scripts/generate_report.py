"""Deterministic report generation. No network, installs, model calls or format fallback."""
import argparse
import html
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unicodedata
from report_schema import read_report, parse_formats, labels, statement_text, cell_text, source_text


def h(value):
    return html.escape(value, quote=True)


def md(value):
    # Preserve literal text in Markdown; no user-authored markup is executable.
    return ''.join('\\' + char if char in '\\`*_{}[]<>()#+.!|~' else char for char in value).replace('\n', '<br>')


def html_document(report):
    out = [f'<!doctype html><html lang="{report["language"]}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{h(report["title"])}</title><style>body{{font-family:"Noto Sans SC",sans-serif;color:#182328;background:#fff;margin:0;line-height:1.7}}main{{max-width:900px;margin:auto;padding:48px 28px}}h1{{font-size:2rem;line-height:1.25}}h2{{margin-top:2rem}}p,li{{overflow-wrap:anywhere}}.table-wrap{{overflow-x:auto}}table{{border-collapse:collapse;width:100%;font-size:.9rem}}th,td{{border:1px solid #cbd5d8;text-align:left;vertical-align:top;padding:10px;min-width:100px;overflow-wrap:anywhere}}th{{background:#eef3f3}}.sources{{font-size:.85rem}}@media(max-width:600px){{main{{padding:24px 16px}}h1{{font-size:1.65rem}}}}@media print{{main{{padding:0}}tr{{break-inside:avoid}}h2{{break-after:avoid}}}}</style></head><body><main>', f'<h1>{h(report["title"])}</h1>', f'<h2>{labels(report)["summary"]}</h2>']
    out += [f'<p>{h(statement_text(item, report))}</p>' for item in report['summary']]
    for section in report['sections']:
        out.append(f'<section id="{section["id"]}"><h2>{h(section["heading"])}</h2>')
        for block in section['blocks']:
            if block['type'] == 'paragraph':
                out.append(f'<p>{h(statement_text(block, report))}</p>')
            elif block['type'] == 'bullets':
                out.append('<ul>' + ''.join(f'<li>{h(statement_text(item, report))}</li>' for item in block['items']) + '</ul>')
            else:
                out.append('<div class="table-wrap"><table><thead><tr>' + ''.join(f'<th scope="col">{h(c)}</th>' for c in block['columns']) + '</tr></thead><tbody>')
                out += ['<tr>' + ''.join(f'<td>{h(cell_text(c, report))}</td>' for c in row) + '</tr>' for row in block['rows']]
                out.append('</tbody></table></div>')
        out.append('</section>')
    out.append(f'<section class="sources"><h2>{labels(report)["sources"]}</h2>')
    out += [f'<p>{h(source_text(source))}</p>' for source in report['sources']]
    out.append('</section></main></body></html>')
    return '\n'.join(out)


def markdown_document(report):
    out = ['# ' + md(report['title']), '## ' + labels(report)['summary']]
    out += [md(statement_text(item, report)) for item in report['summary']]
    for section in report['sections']:
        out.append('## ' + md(section['heading']))
        for block in section['blocks']:
            if block['type'] == 'paragraph':
                out.append(md(statement_text(block, report)))
            elif block['type'] == 'bullets':
                out.append('\n'.join('- ' + md(statement_text(item, report)) for item in block['items']))
            else:
                out.append('| ' + ' | '.join(md(c) for c in block['columns']) + ' |\n| ' + ' | '.join('---' for _ in block['columns']) + ' |\n' + '\n'.join('| ' + ' | '.join(md(cell_text(c, report)) for c in row) + ' |' for row in block['rows']))
    out.append('## ' + labels(report)['sources'])
    out += [md(source_text(source)) for source in report['sources']]
    return '\n\n'.join(out) + '\n'


def estimated_cell_width_mm(value, font_points=10.5):
    """Conservative text estimate, including both cell margins and slack.

    CJK wide/full-width glyphs use one em; Latin and punctuation use .55 em.
    This controls layout only, and never changes content or verification.
    """
    lines = value.splitlines() or ['']
    ems = max(sum(1.0 if unicodedata.east_asian_width(char) in ('W', 'F') else .55 for char in line) for line in lines)
    return ems * font_points * 25.4 / 72 + (180 / 20) * 25.4 / 72 + 1.5


def table_column_widths_mm(block, report, available=170.0):
    natural = [max(estimated_cell_width_mm(block['columns'][index], 11), *(estimated_cell_width_mm(cell_text(row[index], report)) for row in block['rows'])) for index in range(len(block['columns']))]
    short = [width <= 55.0 for width in natural]
    # Keep short cells, including source/status labels, on one line whenever
    # possible. Allocate remaining width to long prose columns independently.
    widths = [max(18.0, width) if is_short else max(36.0, estimated_cell_width_mm(block['columns'][index], 11)) for index, (width, is_short) in enumerate(zip(natural, short))]
    if sum(widths) > available:
        raise ValueError('Table columns exceed readable portrait DOCX width; split the table rather than shrinking or dropping content')
    remaining = available - sum(widths)
    flexible = [index for index, is_short in enumerate(short) if not is_short]
    if flexible:
        demand = sum(natural[index] for index in flexible)
        for index in flexible:
            widths[index] += remaining * natural[index] / demand
    else:
        for index in range(len(widths)):
            widths[index] += remaining / len(widths)
    return widths


def strip_paragraph_borders(element):
    """Remove template paragraph rules without changing text or run styling."""
    namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    properties = element.find(namespace + 'pPr')
    if properties is not None:
        for border in list(properties.findall(namespace + 'pBdr')):
            properties.remove(border)


def write_docx(report, destination):
    from docx import Document
    from docx.shared import Inches, Mm, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
    style = doc.styles['Normal']
    strip_paragraph_borders(style.element)
    style.font.name = 'Noto Sans'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.25
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    for name in ['Title', 'Heading 1', 'Heading 2']:
        heading = doc.styles[name]
        strip_paragraph_borders(heading.element)
        heading.font.name = 'Noto Sans'
        heading.font.size = Pt({'Title': 22, 'Heading 1': 15, 'Heading 2': 12}[name])
        heading.font.color.rgb = RGBColor(0, 0, 0)
        heading.font.underline = False
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(16 if name != 'Title' else 0)
        heading.paragraph_format.space_after = Pt(8)
        color = heading.element.get_or_add_rPr().find(qn('w:color'))
        if color is not None:
            for attribute in ('themeColor', 'themeTint', 'themeShade'):
                color.attrib.pop(qn('w:' + attribute), None)
        heading.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK SC')
    title_paragraph = doc.add_heading(report['title'], 0)
    strip_paragraph_borders(title_paragraph._p)
    doc.add_heading(labels(report)['summary'], 1)
    for item in report['summary']:
        doc.add_paragraph(statement_text(item, report))
    for section_data in report['sections']:
        doc.add_heading(section_data['heading'], 1)
        for block in section_data['blocks']:
            if block['type'] == 'paragraph':
                doc.add_paragraph(statement_text(block, report))
            elif block['type'] == 'bullets':
                for item in block['items']:
                    doc.add_paragraph(statement_text(item, report), style='List Bullet')
            else:
                table = doc.add_table(rows=1, cols=len(block['columns']))
                table.style = 'Table Grid'
                table.autofit = False
                # Protect short cells with their actual citations/status text;
                # long prose receives the remaining width and may wrap alone.
                widths = [Mm(width) for width in table_column_widths_mm(block, report)]
                for column, width in zip(table.columns, widths):
                    column.width = width
                borders = OxmlElement('w:tblBorders')
                for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                    border = OxmlElement('w:' + side)
                    border.set(qn('w:val'), 'single')
                    border.set(qn('w:sz'), '4')
                    border.set(qn('w:color'), 'D9D9D9')
                    borders.append(border)
                table._tbl.tblPr.append(borders)
                header = table.rows[0]
                repeat = OxmlElement('w:tblHeader')
                header._tr.get_or_add_trPr().append(repeat)
                for cell, value in zip(header.cells, block['columns']):
                    cell.text = value
                for row in block['rows']:
                    cells = table.add_row().cells
                    for target, cell in zip(cells, row):
                        target.text = cell_text(cell, report)
                    no_split = OxmlElement('w:cantSplit')
                    table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                for row_index, row in enumerate(table.rows):
                    for cell, width in zip(row.cells, widths):
                        cell.width = width
                        properties = cell._tc.get_or_add_tcPr()
                        margins = OxmlElement('w:tcMar')
                        for side in ('top', 'left', 'bottom', 'right'):
                            margin = OxmlElement('w:' + side)
                            margin.set(qn('w:w'), '90')
                            margin.set(qn('w:type'), 'dxa')
                            margins.append(margin)
                        properties.append(margins)
                        if row_index == 0:
                            shading = OxmlElement('w:shd')
                            shading.set(qn('w:fill'), 'F0F0F0')
                            properties.append(shading)
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.space_after = Pt(3)
                            paragraph.paragraph_format.line_spacing = 1.15
                            for run in paragraph.runs:
                                run.font.size = Pt(10.5)
                                run.font.color.rgb = RGBColor(0, 0, 0)
                                run.font.underline = False
                                run.bold = row_index == 0
    doc.add_heading(labels(report)['sources'], 1)
    for source in report['sources']:
        doc.add_paragraph(source_text(source))
    doc.save(destination)


def preflight(formats):
    if 'docx' in formats or 'pdf' in formats:
        try:
            import docx  # noqa: F401
        except ImportError as error:
            raise RuntimeError('python-docx is required; use the provisioned sandbox, do not install/fallback') from error
        if not shutil.which('fc-match'):
            raise RuntimeError('fc-match is required to verify provisioned document fonts')
        for family in ('Noto Sans', 'Noto Sans CJK SC'):
            result = subprocess.run(['fc-match', '--format=%{family}', family], capture_output=True, text=True, timeout=15, check=True)
            resolved_families = {name.strip().casefold() for name in result.stdout.split(',')}
            if family.casefold() not in resolved_families:
                raise RuntimeError(f'Required font is unavailable: {family}')
    if 'pdf' in formats:
        for command in ('soffice', 'pdftotext', 'pdftoppm'):
            if not shutil.which(command):
                raise RuntimeError(f'{command} is required for the requested PDF path')


def generate(source, output, formats):
    report = read_report(source)
    preflight(formats)
    output = Path(output)
    # mkdir(exist_ok=False) is the contract: never mix outputs from another version.
    output.mkdir(parents=True, exist_ok=False)
    (output / 'report.source.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (output / 'draft.html').write_text(html_document(report), encoding='utf-8')
    if 'md' in formats:
        (output / 'report.md').write_text(markdown_document(report), encoding='utf-8')
    if 'docx' in formats:
        write_docx(report, output / 'report.docx')
    if 'pdf' in formats:
        with tempfile.TemporaryDirectory(prefix='sourceweft-report-') as temporary:
            temp = Path(temporary)
            docx = output / 'report.docx' if 'docx' in formats else temp / 'report.docx'
            if 'docx' not in formats:
                write_docx(report, docx)
            pdf_dir = temp / 'converted'
            pdf_dir.mkdir()
            result = subprocess.run(['soffice', f'-env:UserInstallation={(temp / "profile").as_uri()}', '--headless', '--convert-to', 'pdf', '--outdir', str(pdf_dir), str(docx.resolve())], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=90, check=False)
            converted = pdf_dir / 'report.pdf'
            if result.returncode != 0 or not converted.is_file() or converted.stat().st_size == 0:
                raise RuntimeError('LibreOffice PDF conversion failed; no alternate converter or stale PDF is accepted')
            shutil.move(str(converted), output / 'report.pdf')
    # Generation is not QA. Never emit a passed result here.
    return {'status': 'generated_pending_qa', 'outputDirectory': str(output.resolve()), 'formats': formats, 'htmlDraft': str((output / 'draft.html').resolve()), 'privateSource': str((output / 'report.source.json').resolve())}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', required=True)
    parser.add_argument('--output-dir', required=True)
    parser.add_argument('--formats', default='html,md')
    args = parser.parse_args()
    print(json.dumps(generate(args.source, args.output_dir, parse_formats(args.formats)), ensure_ascii=False))


if __name__ == '__main__':
    main()
